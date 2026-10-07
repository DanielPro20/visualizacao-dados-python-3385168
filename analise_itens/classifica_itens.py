# Classifica os itens de faturamento (descitem) em conceitos regulatórios da Anatel.
# Uso: python classifica_itens.py item_1.xlsx itens_conceitos_anatel.xlsx
import re, sys, unicodedata, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
RGST = " Desde 30/10/2025 as regras gerais estão consolidadas no RGST (Res. 777/2025) e as definições no Glossário Anatel (Res. 779/2025)."

# código -> (categoria, natureza, conceito Anatel, base normativa)
C = {
 "STFC_LOCAL": ("STFC – Chamada local fixo-fixo", "Serviço de telecomunicações (STFC)",
   "Modalidade Local do STFC: comunicação entre pontos fixos determinados situados na mesma Área Local ou em localidades com tratamento local.",
   "Regulamento do STFC (Res. 426/2005)." + RGST),
 "STFC_LOCAL_FM": ("STFC – Chamada local fixo-móvel (VC1)", "Serviço de telecomunicações (STFC)",
   "Chamada originada no STFC (fixo) e terminada em estação móvel do SMP/SME na mesma área de registro; tarifada como VC1 (Valor de Comunicação local).",
   "Regulamento do STFC (Res. 426/2005); regras de VC fixo-móvel." + RGST),
 "STFC_LDN": ("STFC – Longa distância nacional (LDN/DDD) fixo-fixo", "Serviço de telecomunicações (STFC)",
   "Modalidade Longa Distância Nacional do STFC: comunicação entre pontos fixos situados em Áreas Locais distintas do território nacional, sem tratamento local.",
   "Regulamento do STFC (Res. 426/2005)." + RGST),
 "STFC_LDN_FM": ("STFC – LDN fixo-móvel (VC2/VC3)", "Serviço de telecomunicações (STFC)",
   "Chamada de longa distância do fixo para o móvel. VC2: para área de registro distinta com o mesmo 1º dígito do código de área; VC3: para área de registro com 1º dígito diferente.",
   "Regulamento do STFC (Res. 426/2005); conceitos de VC2/VC3." + RGST),
 "STFC_LDI": ("STFC – Longa distância internacional (LDI/DDI)", "Serviço de telecomunicações (STFC)",
   "Modalidade Longa Distância Internacional do STFC: comunicação entre um ponto fixo no Brasil e um ponto em outro país.",
   "Regulamento do STFC (Res. 426/2005)." + RGST),
 "NAO_GEO": ("Numeração não geográfica (0300/0500/0800/0900)", "Serviço de telecomunicações (STFC/SMP) – numeração especial",
   "Códigos Não Geográficos (CNG): números de acesso únicos em todo o território nacional. 0800 – gratuito para quem liga (pago pelo assinante do número); 0300 – quem liga paga o custo de uma chamada local, de qualquer lugar do país, fixo ou móvel; 0500 – doações a entidades de utilidade pública; 0900 – serviços com tarifação diferenciada (valor adicionado ao custo da chamada).",
   "Regulamento de Numeração dos Serviços de Telecomunicações (Res. 749/2022, em vigor desde 03/10/2022); antes, Res. 388/2004 (0300)." + RGST),
 "STFC_PLANO": ("STFC – Plano / assinatura / franquia de telefonia fixa", "Serviço de telecomunicações (STFC)",
   "Valor mensal de plano de serviço do STFC (Plano Básico ou Plano Alternativo de Serviço – PAS), incluindo assinatura e franquia de minutos (ex.: ilimitado local/Brasil).",
   "Regulamento do STFC (Res. 426/2005) – planos de serviço; RGC (Res. 632/2014, substituído pela Res. 765/2023) – oferta e franquia." + RGST),
 "STFC_ACESSO": ("STFC – Acesso / tronco de voz (E1, SIP, interface comutada)", "Serviço de telecomunicações (STFC)",
   "Acesso do assinante ao STFC (linha, tronco digital E1/R2/ISDN-PRI ou tronco SIP) pelo qual são cursadas as chamadas.",
   "Regulamento do STFC (Res. 426/2005)." + RGST),
 "STFC_FACIL": ("Facilidades / comodidades de voz", "Facilidade suplementar do serviço de voz (comodidade)",
   "Funcionalidades adicionais vinculadas ao acesso de voz ou ao número não geográfico (redirecionamento, grupo de busca, URA, agendamento, restrição de área, mensagens, bloqueio etc.). Pela Anatel são comodidades/utilidades adicionais ao serviço; quando não dependem da rede de voz, enquadram-se como SVA.",
   "Regulamento do STFC (Res. 426/2005); RGC; LGT art. 61 (SVA) quando aplicável." + RGST),
 "AUTENT": ("Autenticação de chamadas (STIR/SHAKEN – Origem Verificada)", "Facilidade de voz / exigência regulatória",
   "Assinatura e verificação da identidade de quem origina a chamada (padrão STIR/SHAKEN), adotado pela Anatel no programa ‘Origem Verificada’ contra fraudes e chamadas abusivas.",
   "Atos da Anatel sobre autenticação de chamadas (Grupo de Trabalho de Autenticação de Chamadas)." + RGST),
 "SMP": ("SMP – Telefonia móvel (voz/SMS)", "Serviço de telecomunicações (SMP)",
   "Serviço Móvel Pessoal: serviço móvel terrestre de interesse coletivo que possibilita a comunicação entre estações móveis e destas para outras estações. Tarifação por VC1 (local), VC2 e VC3 (longa distância).",
   "Regulamento do SMP (Res. 477/2007)." + RGST),
 "SMP_ROAM": ("SMP – Roaming nacional", "Serviço de telecomunicações (SMP)",
   "Chamada feita/recebida pelo usuário quando está fora da sua área de registro, atendido como visitante; pode haver adicional por chamada (AD).",
   "Regulamento do SMP (Res. 477/2007); regras de roaming." + RGST),
 "SME": ("SME – Serviço Móvel Especializado (trunking)", "Serviço de telecomunicações (SME)",
   "Serviço móvel de interesse coletivo destinado a pessoas jurídicas ou grupos, com comunicação em grupo/despacho (trunking).",
   "Regulamento do SME." + RGST),
 "SMGS": ("Telefonia móvel por satélite (Iridium)", "Serviço de telecomunicações (SMGS)",
   "Serviço Móvel Global por Satélite: comunicação móvel por meio de constelação de satélites não geoestacionários.",
   "Regulamentação de SMGS / satélites (Res. 748/2021)." + RGST),
 "SMS": ("Mensagens de texto (SMS/torpedo)", "Serviço de telecomunicações (SMP)",
   "Envio de mensagens curtas de texto pela rede móvel, faceta do SMP; serviços de plataforma para disparo corporativo (gestor) são SVA.",
   "Regulamento do SMP (Res. 477/2007); LGT art. 61 para a plataforma." + RGST),
 "MVNO": ("Rede virtual (MVNO)", "Serviço de telecomunicações – atacado (SMP)",
   "Exploração do SMP por meio de rede virtual (credenciado ou autorizado de rede virtual) usando a rede de uma prestadora origem.",
   "Regulamento sobre Exploração de SMP por meio de Rede Virtual (Res. 550/2010)." + RGST),
 "INTERCON": ("Interconexão / remuneração de uso de rede", "Atacado entre prestadoras (interconexão)",
   "Interconexão é a ligação entre redes de telecomunicações funcionalmente compatíveis, para que usuários de uma rede se comuniquem com os de outra. A prestadora que usa a rede de outra paga tarifa/valor de uso: TU-RL (rede local fixa), TU-RIU (rede interurbana), TU-COM (comutação/trânsito), VU-M (rede móvel).",
   "LGT art. 146; Regulamento Geral de Interconexão (Res. 693/2018); valores de referência (Res. 768/2024 e atos de VU-M/EILD)." + RGST),
 "COFAT": ("Repasse / co-faturamento entre prestadoras", "Atacado entre prestadoras",
   "Repasse de receitas e cobrança conjunta: quando o usuário escolhe outra prestadora (ex.: CSP de longa distância) e a cobrança é feita na fatura da prestadora de origem, ou repasse de receita de chamadas de pré-pago.",
   "Regulamento do STFC/SMP – documento de cobrança e co-faturamento; RGC." + RGST),
 "SCM": ("SCM – Internet / banda larga / IP", "Serviço de telecomunicações (SCM)",
   "Serviço de Comunicação Multimídia: serviço fixo de interesse coletivo, em regime privado, que oferece capacidade de transmissão, emissão e recepção de informações multimídia, inclusive conexão à internet, por quaisquer meios.",
   "Regulamento do SCM (Res. 614/2013)." + RGST),
 "SCM_DED": ("SCM – Circuito / linha dedicada / rede corporativa (MPLS, VPN, Frame Relay)", "Serviço de telecomunicações (SCM)",
   "Capacidade de transmissão dedicada ponto a ponto ou multiponto (linha privativa, MPLS, VPN) para o cliente final, prestada como SCM.",
   "Regulamento do SCM (Res. 614/2013)." + RGST),
 "SCM_SAT": ("SCM via satélite (VSAT / IP SAT)", "Serviço de telecomunicações (SCM)",
   "Acesso de dados/internet prestado por satélite (estação VSAT do cliente e hub da prestadora), enquadrado como SCM; usa capacidade de satélite provida conforme regulamento de satélites.",
   "Regulamento do SCM (Res. 614/2013); Regulamento Geral de Satélites (Res. 748/2021)." + RGST),
 "EILD": ("EILD / cessão de meios e capacidade entre prestadoras", "Atacado entre prestadoras",
   "Exploração Industrial: uma prestadora fornece a outra, em regime de exploração industrial, linha dedicada, capacidade ou meios de rede para que esta preste o seu serviço.",
   "Regulamento de EILD (Res. 590/2012); LGT art. 155." + RGST),
 "INFRA": ("Compartilhamento de infraestrutura / fibra apagada / dutos", "Atacado – infraestrutura passiva",
   "Cessão ou compartilhamento de infraestrutura (torres, dutos, postes, espaço, fibra apagada) entre prestadoras. Infraestrutura passiva e fibra apagada não constituem serviço de telecomunicações em si.",
   "Regulamento de Compartilhamento de Infraestrutura (Res. 274/2001, depois Res. 683/2017)." + RGST),
 "RAN": ("Compartilhamento de rede de acesso móvel (RAN sharing)", "Atacado entre prestadoras",
   "Uso compartilhado de elementos da rede de acesso de rádio (ERBs) entre prestadoras móveis, mediante acordo e anuência da Anatel quando exigida.",
   "Regras de compartilhamento de rede/espectro." + RGST),
 "SAT": ("Capacidade de satélite / segmento espacial", "Serviço de telecomunicações – provimento de capacidade",
   "Provimento de capacidade espacial (transponder, segmento espacial) de satélite para transmissão de sinais (dados, vídeo, TV), inclusive uso temporário de frequências.",
   "Regulamento Geral de Satélites (Res. 748/2021)." + RGST),
 "VIDEO": ("Transmissão de sinais de TV / vídeo (contribuição, uplink, eventos)", "Serviço de telecomunicações – transporte de sinais",
   "Transporte de sinais de vídeo de um ponto a outro (uplink, eventos ao vivo) para emissoras/programadoras, usando capacidade terrestre ou de satélite.",
   "Regulamento Geral de Satélites (Res. 748/2021); SCM/SLP conforme o caso." + RGST),
 "SEAC": ("SeAC – TV por assinatura / distribuição de canais", "Serviço de telecomunicações (SeAC)",
   "Serviço de Acesso Condicionado: serviço de interesse coletivo, em regime privado, cuja recepção depende de contratação remunerada, destinado à distribuição de conteúdo audiovisual em pacotes, canais e canais de distribuição obrigatória.",
   "Lei 12.485/2011; Regulamento do SeAC (Res. 581/2012)." + RGST),
 "SLP": ("SLP – Rede privativa", "Serviço de telecomunicações (SLP)",
   "Serviço Limitado Privado: serviço de interesse restrito, em regime privado, para uso próprio da prestadora ou de grupos determinados de usuários (ex.: redes privativas LTE/5G).",
   "Regulamento do SLP (Res. 617/2013), substituído pelo RGST (Res. 777/2025)."),
 "SVA": ("SVA – Serviço de valor adicionado", "Não é serviço de telecomunicações (SVA)",
   "Atividade que acrescenta a um serviço de telecomunicações que lhe dá suporte novas utilidades de acesso, armazenamento, apresentação, movimentação ou recuperação de informações (ex.: e-mail, caixa postal, hospedagem, proteção anti-DDoS, filtragem web). O provedor de SVA é usuário do serviço de telecomunicações.",
   "Lei Geral de Telecomunicações (Lei 9.472/1997), art. 61."),
 "TI": ("TI / software / nuvem (não telecom)", "Não é serviço de telecomunicações",
   "Licenças de software, computação em nuvem, servidores virtuais, segurança de endpoint, contact center, serviços gerenciados e consultoria. Fora da regulação da Anatel; não dependem de telecomunicação para existir como utilidade (no máximo, SVA quando agregados à conexão).",
   "Fora do escopo da Anatel (LGT art. 61 apenas por analogia, quando agregado à conectividade)."),
 "DADOS": ("Serviços de dados / APIs de rede (não telecom)", "Não é serviço de telecomunicações",
   "Consultas e validações baseadas em dados da operadora (score, validação de cadastro, alerta de troca de chip, geolocalização agregada). Sujeitos à LGPD; não são serviço de telecomunicações.",
   "Fora do escopo regulatório de serviço da Anatel; Lei 13.709/2018 (LGPD)."),
 "CONTEUDO": ("Conteúdo / licenciamento / publicidade (não telecom)", "Não é serviço de telecomunicações",
   "Licenciamento e distribuição de conteúdo (música, vídeo), inserção de canais e mídia/publicidade. Atividade de conteúdo, regulada pela Ancine quando audiovisual, não pela Anatel.",
   "Lei 12.485/2011 (separação entre conteúdo e distribuição); fora do escopo da Anatel."),
 "EQUIP": ("Locação / venda / manutenção de equipamentos", "Não é serviço de telecomunicações",
   "Cessão onerosa de bens (modem, roteador, PABX, VSAT, telefones IP, gateways ATA/FXS, notebooks) e sua manutenção. Locação de bens não é serviço de telecomunicações.",
   "Fora do escopo de serviço da Anatel (equipamentos sujeitos apenas a certificação/homologação)."),
 "ACESS": ("Instalação, ativação, mudança e serviços acessórios", "Serviço acessório / habilitação",
   "Valores pela habilitação, instalação, ativação, mudança de endereço ou de velocidade, e visitas técnicas. O RGC exige que esses valores sejam informados na oferta.",
   "RGC (Res. 632/2014, substituído pela Res. 765/2023)." + RGST),
 "DESC": ("Desconto / abatimento", "Ajuste de faturamento",
   "Redução no valor cobrado (promocional, contratual, por volume ou por plano). Quando ligado a uma modalidade (LDN, LDI, local), segue o conceito dela.",
   "RGC (Res. 632/2014, substituído pela Res. 765/2023) – transparência da oferta."),
 "RESSARC": ("Crédito / desconto por interrupção do serviço", "Ajuste de faturamento – obrigação regulatória",
   "Desconto proporcional devido ao consumidor quando o serviço fica interrompido; a Anatel obriga o ressarcimento automático ou mediante reclamação.",
   "RGC (Res. 632/2014, substituído pela Res. 765/2023) – ressarcimento por interrupção."),
 "FRANQ": ("Franquia não utilizada / complemento / compromisso mínimo", "Ajuste de faturamento",
   "Valores ligados à franquia contratada: crédito ou desconto de franquia não consumida, cobrança de complemento até o compromisso mínimo, bônus de minutos.",
   "RGC (Res. 632/2014, substituído pela Res. 765/2023) – franquia e planos."),
 "FIN": ("Multa, juros, parcelamento e encargos", "Não é serviço (encargo financeiro/contratual)",
   "Encargos por atraso (multa, juros, correção), acordos de parcelamento, multa por quebra de fidelidade e penalidades contratuais. O RGC limita e disciplina multa de fidelização e encargos de atraso.",
   "RGC (Res. 632/2014, substituído pela Res. 765/2023) – fidelização e cobrança; Código Civil/CDC."),
 "OUTRAS": ("Outras receitas (imóveis, marca, comissão, cobrança)", "Não é serviço de telecomunicações",
   "Receitas não ligadas à prestação do serviço: aluguel de imóveis/áreas, uso de marca, comissões, prestação de serviço de cobrança, receitas complementares.",
   "Fora do escopo de serviço da Anatel."),
 "INTERCEP": ("Interceptação legal", "Obrigação legal",
   "Atendimento a ordens judiciais de interceptação de comunicações, obrigação das prestadoras.",
   "Lei 9.296/1996; obrigações de suporte à interceptação nos regulamentos de serviço." + RGST),
 "GENERICO": ("Genérico – não é possível classificar só pela descrição", "A validar",
   "A descrição não permite identificar o serviço (ex.: ‘ASSINATURA’, ‘MENSALIDADE’, ‘PROJETO ESPECIAL’). Precisa consultar o produto/contrato de origem.",
   "—"),
}

def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", s).strip()

def has(t, *words):
    return any(re.search(w, t) for w in words)

def classify(t):
    """Retorna (codigo, confianca, observacao)."""
    obs = []
    if "A COBRAR" in t: obs.append("Chamada a cobrar: o pagamento é de quem recebe a chamada.")
    if "DENTRO DA FRANQUIA" in t: obs.append("Minutos consumidos dentro da franquia do plano.")
    if "INTRAGRUPO" in t: obs.append("Chamada entre linhas do mesmo cliente/grupo.")
    if "HORARIO REDUZIDO" in t or re.search(r"\bH\.?R\b|- HR\b", t): obs.append("Horário reduzido (tarifa menor em noites/fins de semana).")
    if "HORARIO NORMAL" in t or re.search(r"- HN\b", t): obs.append("Horário normal.")
    if "OFFNET" in t: obs.append("Destino em rede de outra prestadora (off-net).")
    if "PAS 0" in t: obs.append("PAS = Plano Alternativo de Serviço.")
    o = " ".join(obs)

    # --- ajustes e financeiro primeiro
    if has(t, r"INTERRUP"): return "RESSARC", "Alta", o
    if has(t, r"^DESC", r"DESCONTO"):
        mod = ("LDI" if has(t, "DDI", "LDI", "INTERNACION") else "LDN" if has(t, "DDD", "LDN", "INTERURB", "INTER REGIONAL", "INTRA REGIONAL")
               else "local" if has(t, "LOCAL") else "")
        return "DESC", "Alta", (o + (f" Desconto sobre chamadas {mod}." if mod else "")).strip()
    if has(t, r"FRANQUIA NAO UTILIZADA", r"COMPLEMENTO COMPROMISSO", r"COMPLEMENTO DE FRANQUIA", r"BONUS \d+ MINUTOS", r"VALOR MINIMO DO PLANO", r"UTILIZACAO FRANQUIA"):
        return "FRANQ", "Alta", o
    if has(t, r"SCORE", r"CLARO VALIDA", r"VERIFICACAO DE CPF"): return "DADOS", "Alta", o
    if has(t, r"\bMULTA", r"JUROS", r"ENCARGOS POR ATRASO", r"ATRASO DE PAGAMENTO", r"PARCELAMENTO", r"FIDELIDADE", r"PARCELA DIVIDA",
           r"REFERENTE C\.P\.S", r"DESATIVACAO", r"RETENCAO CONTRATUAL", r"DEVOLUCAO DE VALORES", r"CREDITO", r"PAGAMENTO EM DUPLICIDADE",
           r"COBRANCA TERMINAL INATIVO", r"^CB[AB] - COBRANCA"):
        return "FIN", "Alta", o
    if has(t, r"^AJUSTE"): return "FIN", "Média", "Ajuste de assinatura/serviço (valor de correção na fatura). " + o
    if has(t, r"INTERCEPTACAO"): return "INTERCEP", "Alta", o
    if has(t, r"STIR SHAKEN", r"AUTENTICACAO E IDENTIFICACAO"): return "AUTENT", "Alta", o
    if has(t, r"NAO TELECOM"): return "TI", "Alta", "A própria descrição diz que não é telecom. " + o

    # --- numeração não geográfica
    if has(t, r"NAO GEOGRAFICO", r"\b0300\b", r"\b0500\b", r"\b0800\b", r"\b0900\b", r"SERVICO 900", r"NUMERO INTERNACIONAL", r"SERVICOS HOSTEADOS"):
        extra = ""
        if "NAO GEOGRAFICO" in t:
            extra = " Usa termos do fixo (LOCAL/LDN): corresponde à origem fixo. Agrupa 0300, 0500 e 0900."
        if "NUMERO INTERNACIONAL" in t: extra = " Número de acesso gratuito internacional (padrão UIFN/0800 internacional)."
        if "HOSTEADOS" in t: extra = " Plataforma de rede inteligente que atende o número não geográfico (mensagens, relatórios)."
        return "NAO_GEO", "Alta", (o + extra).strip()

    # --- interconexão / atacado de voz
    if has(t, r"TU-RL", r"TU-COM", r"TU-RIU", r"ENCAMINHAMENTO DE TRAFEGO", r"CHAM\.IU", r"TP-IU"):
        return "INTERCON", "Média" if "TP-IU" in t else "Alta", (o + (" TP-IU: provável tráfego interurbano de rede (validar sigla)." if "TP-IU" in t else "")).strip()
    if has(t, r"REPASSE", r"REP REC", r"CO-FATURAMENTO"): return "COFAT", "Alta", o
    if has(t, r"MVNO"): return "MVNO", "Alta", o
    if has(t, r"RAN SHARING"): return "RAN", "Alta", o
    if has(t, r"EXPLORACAO INDUSTRIAL DE INFRA.*SEAC"): return "EILD", "Média", "Exploração industrial para prestadora de SeAC. " + o
    if has(t, r"EXPLORACAO INDUSTRIAL", r"CESSAO DE MEIOS", r"CESSAO MEIOS", r"CESSAO USO DE REDE", r"CESSAO DE CAPACIDADE"):
        return "EILD", "Alta", o
    if has(t, r"ESTACAO MESTRA"): return "SCM_SAT", "Média", "Uso da estação mestra (hub) satelital da prestadora. " + o
    if has(t, r"FIBRA APAGADA", r"FIBRA OTICA - DUTOS", r"COMPARTILHAMENTO DE INFRA", r"INFRAESTRUTURA ADMINISTRATIV", r"CESSAO DE ESPACO PERMUTA", r"MANUTENCAO DE FIBRAS", r"APAGADA"):
        return "INFRA", "Alta", o
    if has(t, r"^CESSAO - ", r"^CESSAO -", r"CESSAO - CLARO", r"CESSAO - COMCEL", r"CESSAO - AMERICA"):
        return "EILD", "Média", "Cessão a outra empresa do grupo América Móvil (capacidade/meios entre prestadoras de países diferentes). " + o
    if has(t, r"SEGMENTO ESPACIAL", r"SEGMENTOS ESPACIAIS", r"CAPACIDADEDE SATELITE", r"FREQUENCIAS - STAR ONE", r"CESSAO SEGMENTO"):
        return "SAT", "Alta", o
    if has(t, r"UPLINK TV", r"TRANSMISSAO DE EVENTOS", r"EVENTOS AO VIVO", r"SINAIS DE VIDEO") and not has(t, "LOCACAO"):
        return "VIDEO", "Média", o
    if has(t, r"CANAL PERMANENTE"): return "SAT", "Média", "Cessão de espaço em canal permanente (capacidade contínua, provavelmente satelital). " + o
    if has(t, r"DISTRIBUICAO TV"): return "SEAC", "Média", o

    # --- conteúdo, dados, outras receitas
    if has(t, r"CLAROMUSICA", r"CONTENT LICENSING", r"LICENCIAMENTO DE CONTEUDO", r"INSERCAO DA CGTN", r"MIDIA VIA CLARO TV", r"FATURAMENTO REF PLATAFORMA"):
        return "CONTEUDO", "Alta", o
    if has(t, r"CLARO VALIDA", r"CLARO SCORE", r"SCORE DE CREDITO", r"GEODATA", r"NUMBER VERIFICATION", r"SIM SWAP", r"VERIFICACAO DE CPF"):
        return "DADOS", "Alta", o
    if has(t, r"ALUGUEL DE BENS", r"LOCACAO DE AREA", r"UTIL\. DE AREA", r"USO DA MARCA", r"COMISSAO", r"COBRANCA TELEORIENTACAO",
           r"RECEITA COMPLEMENTAR", r"PRESTACAO DE SERVICO DE CO-FATURAMENTO"):
        return "OUTRAS", "Alta", o

    # --- equipamentos
    if has(t, r"ALUG", r"LOCACAO", r"MANUTENCAO", r"^ATA-", r"^GW-", r"^GW ", r"TELECOM.?CPE", r"^TGER", r"^TARS", r"^TBAS", r"ROTEADOR EBT",
           r"^3TEL", r"DESKTOP 8", r"GOOGLE MEET HARDWARE", r"MENSALIDADE PABX"):
        n = "TBAS = aparelho telefônico (não é a ‘assinatura básica’). " if t.startswith("TBAS") else ""
        return "EQUIP", "Alta", (n + o).strip()

    # --- serviços acessórios
    if has(t, r"INSTALA", r"ATIVACAO", r"HABILITACAO", r"MUDANCA DE ENDERECO", r"MUDANCA DE VELOCIDADE", r"MUDANCA ENDERECO",
           r"AGENDAMEN", r"VRE1", r"CREDENCIAMENTO", r"TAXA DE BLOQUEIO"):
        if has(t, r"AGENDAMEN") and not has(t, "INSTALA"):
            return "STFC_FACIL", "Média", "Facilidade de agendamento (programação de horário/data de atendimento do número). " + o
        if has(t, r"MICROSOFT", r"OMNI"): return "TI", "Alta", "Taxa de ativação de licença/serviço de TI. " + o
        return "ACESS", "Alta", o

    # --- TI / SVA
    ti = [r"MICROSOFT", r"OFFICE", r"O365", r"M365", r"EXCHANGE", r"SHAREPOINT", r"ONEDRIVE", r"TEAMS", r"SKYPE", r"POWER (BI|APPS|AUTOMATE)",
          r"DYNAMICS", r"DATAVERSE", r"VISIO", r"PROJECT", r"PLANNER", r"COPILOT", r"AZURE", r"ENTRA ID", r"INTUNE", r"PURVIEW", r"DEFENDER",
          r"WINDOWS", r"REDHAT", r"SQL", r"GOOGLE WORKSPACE", r"CLOUD", r"AWS", r"MULTICLOUD", r"VCPU", r"MEMORIA", r"INSTANCIA",
          r"DATA CENTER VIRTUAL", r"DRIVE", r"NUVEM", r"BACKUP", r"ARMAZ", r"STORAGE", r"EDR", r"SOPHOS", r"KASPERSK", r"MDM", r"MAAS360", r"MOBILE DEVICE",
          r"OMNI", r"CONTACT CENTER", r"AMAZON CONNECT", r"DE PAS\b", r"/ ?QTD", r"ATEND", r"\bBOT\b", r"\bURA COGN", r"TELEMEDICINA", r"LICENC",
          r"CONSULTORIA", r"GERENCIA", r"GESTAO", r"OUTSOURCING", r"SUPORTE TECNICO", r"UST\b", r"PROCESSAMENTO DE DADOS", r"SIEM",
          r"SECURITY ANALYTICS", r"VULNERABILIDADES", r"CYBER", r"^PT ", r"RETESTE", r"SMART VIEW", r"TELEPRESENCA", r"WEBEX",
          r"SOFTPHONE", r"^SOFT", r"MARKETPLACE", r"FORTITOKEN", r"FLEX SECURITY", r"BLD SEGURO", r"TI -", r"TI-", r"MENSALIDADE TI",
          r"\bGRC\b", r"MONITORAMENTO DE SATELITE", r"MONITORACAO DE TRAFEGO", r"SLA VIEW", r"PROJETO NOSSA TV", r"MSP", r"FINOPS",
          r"GFW\d", r"ADMINISTRACAO DE INFRAESTRUTURA", r"FULL ", r"FULL/", r"SERVICOS DIGITAIS", r"CONECTA COM", r"COMBO CONECTA", r"PABX VIRTUAL",
          r"PACOTE (AVANCADO|BASICO)", r"BUSINESS SECURITY", r"SERVICO B\. SECURITY", r"BHOST", r"AGENTE PARA SERVIDOR", r"MENSALIDADE TI-HCC",
          r"MENSALIDADE SCV2", r"EMVIA", r"MENSALIDADE TP", r"^TP - ", r"PROCESSAMENTO PARCEIRO", r"DESKTOP", r"TI - GERENCIA"]
    sva = [r"ANTI-DDOS", r"\bWAF", r"INTERNET SEGURA", r"WEB GATEWAY", r"SECURE WEB", r"WIFI SEGURO", r"WI-FI", r"ACCESS POINT", r"ACESS POINT", r"AC\.INT",
           r"CAIXA POSTAL", r"CAIXA POS", r"CX\.POSTAL", r"E-MAIL", r"CONTA DE E", r"HOSPEDAGEM", r"CONSTRUTOR DE SITES", r"SITE PRONTO",
           r"DOMINIO", r"SITE MOVEL", r"LOAD BALANCE", r"CONEXAO BACKEND", r"TRANSFERENCIA", r"DCV PUBLICO", r"OFFICE DIAL", r"INTERNET SIMPLES",
           r"SENHA", r"MENSAGEM", r"^MENS[\. ]", r"MESG", r"RELATORIO", r"BINA", r"GESTOR CONEXAO TORPEDO", r"SD-?WAN", r"SEGURANCA DE ACESSO",
           r"ACESSO RESTRITO"]
    if has(t, r"^ACESSO RESTRITO$", r"^CONEXAO$", r"^INFRA-ESTRUTURA$", r"SUPPLY"):
        return "GENERICO", "Baixa", ("‘Supply’: fornecimento/ponto de equipamento — validar. " if "SUPPLY" in t else "") + o
    if has(t, r"^AC\.INT"): return "SCM", "Alta", "Acesso à internet (DG1 = degrau 1 de distância; EBT = Embratel). " + o
    if has(t, r"GERENCIA", r"OUTSOURCING"): return "TI", "Alta", "Serviço gerenciado (gestão de rede/equipamentos). " + o
    if has(t, r"SD.?W.?AN") and has(t, r"CONECTIVIDADE"): return "SCM_DED", "Média", "Conectividade SD-WAN (transporte); gestão/licença de SD-WAN é TI/SVA. " + o
    if has(t, *sva) and not has(t, r"^MENSALIDADE ACESSO"):
        n = ""
        if has(t, r"^MENS[\. ]", r"MESG", r"MENSAGEM PADRONIZADA", r"MENSAGEM PERSONALIZADA", r"MENSAGEM (DE )?NAVEGACAO"):
            return "STFC_FACIL", "Média", "Mensagens gravadas/de navegação ou tarifas por tamanho/horário de mensagem (H.N/H.R). Validar se é URA de número não geográfico ou serviço de mensagens de dados. " + o
        if has(t, r"SENHA"): return "STFC_FACIL", "Média", "Senha de acesso/gerente vinculada ao serviço. " + o
        if has(t, r"SD-?WAN"): n = "Gestão/licença/segurança de SD-WAN sobre a conectividade. "
        return "SVA", "Alta", (n + o).strip()
    if has(t, r"SERVICOS DIGITAIS FONE"): return "STFC_FACIL", "Média", "Pacote de serviços digitais da linha fixa (identificador de chamadas, siga-me etc.) — validar. " + o
    if has(t, *ti):
        if has(t, r"TELEFONE IP", r"SOFTPHONE", r"PHONE", r"VOZ", r"PBX", r"RAMAL"):
            return "TI", "Média", "Telefonia IP/PABX em nuvem: a plataforma e o software são TI/SVA; a saída para a rede pública (números e chamadas) é STFC. " + o
        return "TI", "Alta", o
    if has(t, r"MENSALIDADE (PLANO|EMPRESARIAL|EXPRESSO|LINUX)"):
        return "SVA", "Média", "Pela vizinhança na lista (Exchange/O365 ‘Empresarial’, ‘Expresso’), provável plano de hospedagem/e-mail — validar. " + o
    if has(t, r"^ASSINATURA (EMPRESARIAL|INDIVIDUAL)$"): return "GENERICO", "Baixa", o

    # --- facilidades de voz
    if has(t, r"FACILIDADE", r"FAC\.", r"GRUPO DE BUSCA", r"REENCAMINHAMENTO", r"REDIRECI", r"RESTRICAO", r"RETRICAO", r"LIMITACAO CHAMADAS",
           r"MSG MUDAN", r"SELECAO DE ORIGEM", r"ROTA DIRETA", r"ROTEAM ALTERNATIVO", r"DISTRIB\.PERCEN", r"NUMERACAO", r"PACOTE DE FACILIDADES"):
        return "STFC_FACIL", "Alta", o

    # --- satélite / dados
    if has(t, r"IRIDIUM"): return "SMGS", "Alta", o
    if has(t, r"SATELITE", r"\bSAT\b", r"IPSAT", r"IP SAT", r"VSAT", r"VIPSAT", r"ESTACAO MESTRA", r"GESAC", r"INTERNET S AT", r"INTER\. SA T"):
        return "SCM_SAT", "Alta", o
    if has(t, r"CIRCUITO", r"CCTO", r"PRIVATE LINE", r"DEDICATED", r"TRANSPARENT LAN", r"ACESSO DEDICADO", r"MPLS", r"FRAME RELAY", r"VPN",
           r"RENPAC", r"PORTA RENPAC", r"DEGRAU", r"\(DG\d\)", r"ACESSO INTERURBANO", r"REDE PRIVATIVA", r"NET VIRTUAL", r"CANAL LOGICO",
           r"UPLINK DE DADOS", r"TRANSMISSAO DO SERVICO"):
        if has(t, r"REDE PRIVATIVA", r"PROG REDE PRIVATIVA"): return "SLP", "Média", "Rede privativa para uso de grupo determinado (provável SLP). " + o
        if has(t, r"RENPAC"): return "SCM_DED", "Alta", "RENPAC: antiga rede pública de comutação de pacotes da Embratel (X.25), hoje enquadrada como SCM. " + o
        return "SCM_DED", "Alta", o
    if has(t, r"ACESSO", r"INTERNET", r"BANDA LARGA", r"IP FIXO", r"ENDERECO IP", r"^INTERFACE (?!COMUTADA|TELEF|DED\.64|VIPPHONE)", r"^PORTA",
           r"BACKBONE", r"FASTNET", r"BUSINESS LINK", r"PREST.* TELECOMUNICACOES.*(BANDA|MBPS)", r"TELECOM BANDA", r"TELECOM ACESSO",
           r"TRAFEGO COMPL\. 2MBPS", r"CONEXAO$", r"DISPONIBILIDADE - 2M", r"INFRA-ESTRUTURA"):
        conf = "Média" if has(t, r"CONEXAO$", r"DISPONIBILIDADE", r"INFRA-ESTRUTURA", r"ACESSO LOCAL$", r"FACILIDADE DE ACESSO") else "Alta"
        return "SCM", conf, o

    # --- voz STFC / SMP
    if has(t, r"SME\b", r"-SME"): return "SME", "Alta", o
    if has(t, r"TORPEDO"): return "SMS", "Alta", o
    if has(t, r"ROAMING"): return "SMP_ROAM", "Alta", o
    movel_origem = has(t, r"MOVEL[- ]FIXO", r"MOVEL[- ]MOVEL", r"MOVEL MOVEL", r"MOVEL FIXO", r"ORIGEM MOVEL", r"MOVEL INTERURBANO - PLANO", r"MOVEL LOCAL - PLANO",
                       r"MOVEL RESIDENTE", r"MOVEL VISITANTE")
    if has(t, r"TRONCO", r"^TSIP", r"INTERFACE COMUTADA", r"INTERFACE TELEF", r"INTERFACE DED\.64", r"INTERFACE VIPPHONE", r"INTERFACE DEDICADA"):
        return "STFC_ACESSO", "Média" if "DEDICADA" in t else "Alta", o
    if has(t, r"DDI", r"LDI", r"INTERNACION", r"BRASIL DIRETO"):
        if has(t, r"MOVEL FIXO CLARO"): return "SMP", "Alta", "Chamada internacional originada no móvel. " + o
        return "STFC_LDI", "Alta", o
    if movel_origem and not has(t, r"RECEBIDAS A COBRAR", r"A COBRAR DE CELULAR", r"MOVEL-FIXO A COBRAR"):
        return "SMP", "Alta", o
    ld = has(t, r"DDD", r"LDN", r"INTERURBAN", r"INTER REGIONAL", r"INTRA REGIONAL", r"INTER-ESTADUAL", r"INTRA-ESTADUAL", r"NACIONA", r"\bLD\b",
             r"FRANQUIA F-F LDN")
    fm = has(t, r"FIXO[- ]MOVEL", r"FIXO MOVEL", r"P/MOVEL", r"PARA CELULAR", r"PARA REDE MOVEL", r"TELEFONE MOVEL", r"\bF-M\b", r"REDE MOVEL",
             r"A COBRAR DE CELULAR", r"RECEBIDAS A COBRAR DE CELULARES", r"MOVEL-FIXO A COBRAR", r"TRAFEGO MOVEL", r"PARA CELULARES", r"DE CELULAR")
    if not has(t, r"LIGACOES") and has(t, r"^0[25] ?-", r"^02-", r"ILIMITADO", r"ILIM\.", r"NETFONE", r"NET FONE", r"FALE FIXO", r"MENSALIDADE TELEFONIA", r"FRANQUIA DE TELEFONIA",
           r"PACOTE DE FRANQUIA", r"ASSINATURA REG", r"PLANO MUITO MAIS", r"TARIFA FLAT", r"MENSALIDADE (PLANO|EMPRESARIAL|EXPRESSO|LINUX)",
           r"ASSINATURA (EMPRESARIAL|INDIVIDUAL)", r"FRANQUIA F-F"):
        if has(t, r"ACESSO INTERNET"): return "SCM", "Alta", o
        conf = "Média" if has(t, r"MENSALIDADE (PLANO|EMPRESARIAL|EXPRESSO|LINUX)", r"ASSINATURA (EMPRESARIAL|INDIVIDUAL)") else "Alta"
        n = "Descrição genérica de mensalidade; assumido plano de telefonia — validar. " if conf == "Média" else ""
        return "STFC_PLANO", conf, (n + o).strip()
    if ld:
        n = "Chamada destinada: cobrada do assinante de destino (típico de número 0800/serviço de recebimento) — validar. " if "DESTINADA" in t else ""
        return ("STFC_LDN_FM" if fm else "STFC_LDN"), "Média" if n else "Alta", (n + o).strip()
    if has(t, r"LOCAL", r"LOCAIS", r"ON - NET", r"FORA DA AREA DE TARIFACAO", r"CENTRO VIRTUAL", r"DESTINADA"):
        n = ""
        if "DESTINADA" in t: n = "Chamada destinada: cobrada do assinante de destino (típico de número 0800/serviço de recebimento) — validar. "
        if "FORA DA AREA DE TARIFACAO" in t: n = "Chamada local para fora da Área de Tarifação Básica (ATB). "
        return ("STFC_LOCAL_FM" if fm else "STFC_LOCAL"), "Média" if n.startswith("Chamada destinada") else "Alta", (n + o).strip()
    if has(t, r"CHAMADA", r"LIGACOES", r"TRAFEGO"): return "STFC_LOCAL", "Baixa", "Chamada sem modalidade explícita. " + o
    return "GENERICO", "Baixa", o

wb = openpyxl.load_workbook(SRC, data_only=True)
items = [r[0] for r in wb.active.iter_rows(min_row=2, values_only=True) if r and r[0]]
rows = []
for it in items:
    code, conf, obs = classify(norm(str(it)))
    cat, nat, conc, base = C[code]
    rows.append((str(it).strip(), cat, nat, conc, base, conf, obs))

out = openpyxl.Workbook()
hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F4E78")
def sheet(ws, header, data, widths):
    ws.append(header)
    for c in ws[1]: c.font, c.fill, c.alignment = hdr_font, hdr_fill, Alignment(vertical="center", wrap_text=True)
    for r in data: ws.append(list(r))
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions

ws = out.active; ws.title = "Itens"
sheet(ws, ["Item (descitem)", "Categoria", "Natureza regulatória", "Conceito (Anatel)", "Base normativa", "Confiança", "Observação"],
      rows, [55, 38, 30, 70, 55, 11, 55])
cnt = collections.Counter(r[1] for r in rows)
ws2 = out.create_sheet("Conceitos")
sheet(ws2, ["Categoria", "Natureza regulatória", "Conceito (Anatel)", "Base normativa", "Qtd. itens"],
      [(v[0], v[1], v[2], v[3], cnt.get(v[0], 0)) for v in C.values()], [45, 32, 80, 60, 10])
ws3 = out.create_sheet("Resumo")
nat = collections.Counter(r[2] for r in rows); conf = collections.Counter(r[5] for r in rows)
sheet(ws3, ["Agrupamento", "Valor", "Qtd. itens"],
      [("Natureza", k, v) for k, v in nat.most_common()] + [("Confiança", k, v) for k, v in conf.most_common()], [15, 60, 12])
out.save(OUT)
print(len(rows), "itens"); print(conf); 
for k, v in cnt.most_common(): print(f"{v:5} {k}")
